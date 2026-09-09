from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

NUMBER_FORMAT = "0.0000"

_LABEL_COLUMN_WIDTH = 14
_METRIC_COLUMN_WIDTH = 12


def write_grid(
    path: Path | str,
    *,
    models: Sequence[str],
    labels: Sequence[str],
    metrics: Sequence[str],
    cells: Mapping[tuple[str, str], tuple[float, ...]],
    sheet_title: str,
) -> Path:
    path = Path(path)
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title

    _write_header(sheet, models, metrics)
    for offset, label in enumerate(labels):
        _write_row(
            sheet,
            row=3 + offset,
            label=label,
            models=models,
            metrics=metrics,
            cells=cells,
        )
    _set_widths(sheet, models, metrics)
    sheet.freeze_panes = "B3"

    book.save(path)
    return path


def _first_column_of(index: int, metrics: Sequence[str]) -> int:
    return 2 + len(metrics) * index


def _write_header(sheet, models: Sequence[str], metrics: Sequence[str]) -> None:
    corner = sheet.cell(row=1, column=1)
    corner.alignment = Alignment(horizontal="center", vertical="center")
    if models:
        sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    for index, model in enumerate(models):
        first = _first_column_of(index, metrics)
        name = sheet.cell(row=1, column=first, value=model)
        name.font = Font(bold=True)
        name.alignment = Alignment(horizontal="center")
        sheet.merge_cells(
            start_row=1,
            start_column=first,
            end_row=1,
            end_column=first + len(metrics) - 1,
        )
        for offset, metric in enumerate(metrics):
            cell = sheet.cell(row=2, column=first + offset, value=metric)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")


def _write_row(
    sheet,
    *,
    row: int,
    label: str,
    models: Sequence[str],
    metrics: Sequence[str],
    cells: Mapping[tuple[str, str], tuple[float, ...]],
) -> None:
    heading = sheet.cell(row=row, column=1, value=label)
    heading.font = Font(bold=True)

    published = [
        values for model in models if (values := cells.get((model, label))) is not None
    ]
    best = [
        min(values[metric] for values in published) if len(published) > 1 else None
        for metric in range(len(metrics))
    ]

    for index, model in enumerate(models):
        values = cells.get((model, label))
        if values is None:
            continue
        first = _first_column_of(index, metrics)
        for metric, value in enumerate(values):
            cell = sheet.cell(row=row, column=first + metric, value=value)
            cell.number_format = NUMBER_FORMAT
            if value == best[metric]:
                cell.font = Font(bold=True)


def _set_widths(sheet, models: Sequence[str], metrics: Sequence[str]) -> None:
    sheet.column_dimensions["A"].width = _LABEL_COLUMN_WIDTH
    for column in range(2, _first_column_of(len(models), metrics)):
        sheet.column_dimensions[get_column_letter(column)].width = _METRIC_COLUMN_WIDTH
